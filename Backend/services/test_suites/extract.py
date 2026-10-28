"""Deterministic test suite extraction: one JSON block per unique test id, values copied from the cells."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

# Canonical field → header names it may appear under (compared case/space/punctuation-insensitively).
FIELDS: dict[str, tuple[str, ...]] = {
    "test_id": ("TestId", "Test Id", "Test Case Id", "TC Id", "TCID", "Test Case No", "Test Case Number"),
    "area": ("Area", "Domain", "Module"),
    "group_id": ("GroupId", "Group Id"),
    "functional_group": ("Functional Group", "Feature", "Category"),
    "short_description": ("Short description", "Title", "Test Case Name", "Test Name", "Summary"),
    "details": ("Details",),
    "description": ("Description", "Test Description", "Test Case Description", "Objective"),
    "pre_conditions": ("Pre Conditions", "Preconditions", "Pre-requisites", "Prerequisites"),
    "testing_steps": ("Testing Steps", "Test Steps", "Steps", "Procedure"),
    "expected_results": ("Expected Results", "Expected Result", "Expected Outcome"),
    "repetitions": ("Repetitions", "Iterations"),
    "route_dep": ("Route_dep", "Route Dependency"),
    "route_desc": ("Route_desc", "Route Description"),
    "network_dep": ("Network_dep", "Network Dependency"),
    "priority": ("Priority",),
    "context": ("Context",),
    "in_size": ("In size",),
    "notes": ("NOTES", "Remarks", "Comments"),
}
REQUIRED_TEXT = ("testing_steps", "expected_results")
ENUMS = {"route_dep": {"YES", "NO"}, "network_dep": {"YES", "NO"}, "in_size": {"Y", "N"}}
PRIORITY_RE = re.compile(r"^P\d$")
TRUE_FLAGS = {"y", "yes", "x", "true", "✓"}
FALSE_FLAGS = {"n", "no", "false"}
HEADER_SCAN_ROWS = 20
MIN_HEADER_FIELDS = 3
UNGROUPED = "ungrouped"


def _norm_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value).lower())


ALIAS_TO_FIELD = {_norm_header(a): f for f, aliases in FIELDS.items() for a in aliases}


def _clean(value: Any) -> Any:
    """Cell value as JSON: whole floats become ints, dates ISO strings, text trimmed with \\n line breaks."""
    if value is None or isinstance(value, bool):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    text = str(value).replace("\r\n", "\n").replace("\r", "\n").strip()
    return text or None


def _flag(value: Any) -> bool | None:
    """True/False for flag-like cells ("Y", "X", "NO", blank); None when the value is not a flag."""
    if value is None:
        return False
    if isinstance(value, bool):
        return value
    text = str(value).strip().lower()
    if text in TRUE_FLAGS:
        return True
    if text in FALSE_FLAGS:
        return False
    return None


@dataclass
class Column:
    index: int
    letter: str
    header: str
    group: str | None
    field: str | None = None
    membership: tuple[str, str] | None = None
    extra_key: str | None = None

    @property
    def maps_to(self) -> str:
        if self.field:
            return self.field
        if self.membership:
            return f"membership.{self.membership[0]}.{self.membership[1]}"
        return f"extra.{self.extra_key}"


@dataclass
class SheetResult:
    info: dict[str, Any]
    rows: list[tuple[int, dict[str, Any]]] = field(default_factory=list)
    skipped: list[dict[str, Any]] = field(default_factory=list)
    checks: list[dict[str, Any]] = field(default_factory=list)


def _grid(ws) -> list[list[Any]]:
    return [
        list(r)
        for r in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column, values_only=True)
    ]


def _uncached_formulas(formula_ws, grid: list[list[Any]]) -> list[str]:
    """Formula cells whose last calculated value was not saved in the file (they would read as empty)."""
    missing = []
    for r, row in enumerate(
        formula_ws.iter_rows(min_row=1, max_row=len(grid), min_col=1, max_col=len(grid[0]) if grid else 1, values_only=True)
    ):
        for c, v in enumerate(row):
            if isinstance(v, str) and v.startswith("=") and grid[r][c] is None:
                missing.append(f"{get_column_letter(c + 1)}{r + 1}")
    return missing


def _find_header(grid: list[list[Any]]) -> int | None:
    best_count, best_idx = 0, None
    for i, row in enumerate(grid[:HEADER_SCAN_ROWS]):
        fields = {ALIAS_TO_FIELD.get(_norm_header(v)) for v in row if v is not None} - {None}
        if "test_id" in fields and len(fields) > best_count:
            best_count, best_idx = len(fields), i
    return best_idx if best_count >= MIN_HEADER_FIELDS else None


def _group_band(ws, grid: list[list[Any]], header_idx: int) -> dict[int, str]:
    """Labels in the row above the header (e.g. merged "PAT" over S1:V1) per 0-based column."""
    if header_idx == 0:
        return {}
    band_row = header_idx  # 1-based row number of the row above the header
    band: dict[int, str] = {}
    for c, v in enumerate(grid[header_idx - 1]):
        label = _clean(v)
        if label is not None:
            band[c] = str(label)
    for rng in ws.merged_cells.ranges:
        if rng.min_row == band_row and (rng.min_col - 1) in band:
            for c in range(rng.min_col - 1, rng.max_col):
                band[c] = band[rng.min_col - 1]
    return band


def _fill_merged(ws, grid: list[list[Any]], first_data_idx: int) -> None:
    """A merged block below the header (e.g. one precondition for 5 rows) applies to every row it covers."""
    for rng in ws.merged_cells.ranges:
        if rng.min_row - 1 < first_data_idx:
            continue
        value = grid[rng.min_row - 1][rng.min_col - 1]
        for r in range(rng.min_row - 1, rng.max_row):
            for c in range(rng.min_col - 1, rng.max_col):
                grid[r][c] = value


def _plan_columns(header: list[Any], band: dict[int, str], data_rows: list[list[Any]]) -> list[Column]:
    columns: list[Column] = []
    used_fields: set[str] = set()
    for idx, raw in enumerate(header):
        has_data = any(_clean(r[idx]) is not None for r in data_rows if idx < len(r))
        name = _clean(raw)
        if name is None and not has_data:
            continue
        letter = get_column_letter(idx + 1)
        col = Column(index=idx, letter=letter, header=str(name) if name is not None else f"Column {letter}", group=band.get(idx))
        mapped = ALIAS_TO_FIELD.get(_norm_header(col.header)) if name is not None else None
        if mapped and mapped not in used_fields:
            col.field = mapped
            used_fields.add(mapped)
        columns.append(col)

    membership_labels: set[tuple[str, str]] = set()
    for col in columns:
        if col.field:
            continue
        values = [r[col.index] for r in data_rows if col.index < len(r) and _clean(r[col.index]) is not None]
        if values and all(_flag(v) is not None for v in values):
            group = col.group or UNGROUPED
            label = col.header
            if col.group and label.upper().startswith(f"{col.group.upper()}_"):
                label = label[len(col.group) + 1 :]
            if (group, label) in membership_labels:
                label = f"{label} ({col.letter})"
            membership_labels.add((group, label))
            col.membership = (group, label)
        else:
            col.extra_key = f"{col.group} / {col.header}" if col.group else col.header
    return columns


def _row_record(cells: list[Any], columns: list[Column]) -> dict[str, Any]:
    record: dict[str, Any] = {f: None for f in FIELDS}
    membership: dict[str, dict[str, bool]] = {}
    extra: dict[str, Any] = {}
    for col in columns:
        raw = cells[col.index] if col.index < len(cells) else None
        if col.field:
            record[col.field] = _clean(raw)
        elif col.membership:
            group, label = col.membership
            membership.setdefault(group, {})[label] = bool(_flag(raw))
        else:
            value = _clean(raw)
            if value is not None:
                extra[col.extra_key] = value
    record["membership"] = membership
    record["extra"] = extra
    return record


def _extract_sheet(ws, formula_ws) -> SheetResult:
    info: dict[str, Any] = {"name": ws.title, "visible": ws.sheet_state == "visible"}
    if not info["visible"]:
        return SheetResult({**info, "status": "skipped", "reason": "Hidden sheet"})
    grid = _grid(ws)
    header_idx = _find_header(grid)
    if header_idx is None:
        return SheetResult(
            {
                **info,
                "status": "skipped",
                "reason": f"No test case header row (needs a test id column and {MIN_HEADER_FIELDS}+ known columns)",
            }
        )

    uncached = _uncached_formulas(formula_ws, grid)
    band = _group_band(ws, grid, header_idx)
    _fill_merged(ws, grid, header_idx + 1)
    header = grid[header_idx]
    width = max(len(r) for r in grid)
    body = [r + [None] * (width - len(r)) for r in grid[header_idx + 1 :]]
    header = header + [None] * (width - len(header))
    id_col = next(i for i, v in enumerate(header) if ALIAS_TO_FIELD.get(_norm_header(v)) == "test_id")

    test_rows: list[tuple[int, list[Any]]] = []
    totals_rows: list[tuple[int, list[Any]]] = []
    result = SheetResult(info)
    for offset, cells in enumerate(body):
        row_no = header_idx + 2 + offset
        filled = [_clean(v) for v in cells if _clean(v) is not None]
        if not filled:
            continue
        if _clean(cells[id_col]) is not None:
            test_rows.append((row_no, cells))
        elif all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in filled):
            totals_rows.append((row_no, cells))
        else:
            result.skipped.append(
                {"sheet": ws.title, "row": row_no, "reason": "No test id", "preview": [str(v)[:60] for v in filled[:3]]}
            )

    columns = _plan_columns(header, band, [c for _, c in test_rows])
    for row_no, cells in test_rows:
        result.rows.append((row_no, _row_record(cells, columns)))

    for row_no, cells in totals_rows:
        for col in columns:
            if not col.membership:
                continue
            expected = _clean(cells[col.index])
            if not isinstance(expected, (int, float)):
                continue
            actual = sum(1 for _, c in test_rows if _flag(c[col.index]))
            result.checks.append(
                {
                    "name": f"{ws.title}: {f'{col.group} ' if col.group else ''}{col.header} count matches totals row {row_no}",
                    "ok": actual == expected,
                    "detail": f"{actual} rows marked, totals row says {expected}",
                }
            )
    result.checks.append(
        {
            "name": f"{ws.title}: formulas have saved values",
            "ok": not uncached,
            "detail": f"{len(uncached)} formula cells without a saved value"
            + (f" (e.g. {', '.join(uncached[:5])})" if uncached else ""),
        }
    )

    result.info = {
        **info,
        "status": "extracted",
        "header_row": header_idx + 1,
        "group_row": header_idx if band else None,
        "columns": [{"column": c.letter, "header": c.header, "group": c.group, "maps_to": c.maps_to} for c in columns],
        "test_rows": len(test_rows),
        "totals_rows": [r for r, _ in totals_rows],
        "skipped_rows": len(result.skipped),
    }
    return result


def _distinct(values: list[tuple[dict[str, Any], Any]]) -> list[tuple[dict[str, Any], Any]]:
    seen: list[Any] = []
    out = []
    for src, v in values:
        if v not in seen:
            seen.append(v)
            out.append((src, v))
    return out


def _merge(test_id: str, rows: list[tuple[dict[str, Any], dict[str, Any]]]) -> dict[str, Any]:
    """One block per test id: equal values kept once, membership flags OR-ed, differing values listed as conflicts."""
    first = rows[0][1]
    merged: dict[str, Any] = {"test_id": test_id, "sources": [src for src, _ in rows]}
    conflicts: list[dict[str, Any]] = []

    for name in FIELDS:
        if name == "test_id":
            continue
        merged[name] = first[name]
        variants = _distinct([(src, rec[name]) for src, rec in rows])
        if len(variants) > 1:
            conflicts.append(
                {"field": name, "resolution": "first_row", "values": [{**src, "value": v} for src, v in variants]}
            )

    membership: dict[str, dict[str, bool]] = {}
    for group in dict.fromkeys(g for _, rec in rows for g in rec["membership"]):
        labels = dict.fromkeys(l for _, rec in rows for l in rec["membership"].get(group, {}))
        membership[group] = {}
        for label in labels:
            values = [(src, rec["membership"].get(group, {}).get(label, False)) for src, rec in rows]
            membership[group][label] = any(v for _, v in values)
            if len(_distinct(values)) > 1:
                conflicts.append(
                    {
                        "field": f"membership.{group}.{label}",
                        "resolution": "any_row",
                        "values": [{**src, "value": v} for src, v in values],
                    }
                )
    merged["membership"] = membership

    extra: dict[str, Any] = {}
    for key in dict.fromkeys(k for _, rec in rows for k in rec["extra"]):
        variants = _distinct([(src, rec["extra"].get(key)) for src, rec in rows])
        extra[key] = first["extra"].get(key)
        if len(variants) > 1:
            conflicts.append(
                {"field": f"extra.{key}", "resolution": "first_row", "values": [{**src, "value": v} for src, v in variants]}
            )
    merged["extra"] = extra
    merged["conflicts"] = conflicts
    return merged


def _warnings(case: dict[str, Any], present: set[str]) -> list[dict[str, str]]:
    out = []
    for name in REQUIRED_TEXT:
        if name in present and case[name] is None:
            out.append({"field": name, "issue": "Empty"})
    if ({"short_description", "description"} & present) and not (case["short_description"] or case["description"]):
        out.append({"field": "description", "issue": "No short description or description"})
    if "functional_group" in present and case["functional_group"] is None:
        out.append({"field": "functional_group", "issue": "Empty"})
    for name, allowed in ENUMS.items():
        value = case[name]
        if name in present and value is not None and str(value).upper() not in allowed:
            out.append({"field": name, "issue": f"Unexpected value {value!r} (expected {'/'.join(sorted(allowed))})"})
    if "priority" in present and case["priority"] is not None and not PRIORITY_RE.match(str(case["priority"])):
        out.append({"field": "priority", "issue": f"Unexpected value {case['priority']!r} (expected P0-P9)"})
    return out


def extract_workbook(path: Path) -> dict[str, Any]:
    wb = load_workbook(path, data_only=True)
    formula_wb = load_workbook(path, read_only=True)
    try:
        sheet_results = [_extract_sheet(ws, formula_wb[ws.title]) for ws in wb.worksheets]
    finally:
        wb.close()
        formula_wb.close()

    present: set[str] = set()
    by_id: dict[str, list[tuple[dict[str, Any], dict[str, Any]]]] = {}
    for res in sheet_results:
        present |= {c["maps_to"] for c in res.info.get("columns", [])}
        for row_no, record in res.rows:
            src = {"sheet": res.info["name"], "row": row_no}
            by_id.setdefault(str(record["test_id"]), []).append((src, record))

    test_cases = []
    for test_id, rows in by_id.items():
        case = _merge(test_id, rows)
        case["warnings"] = _warnings(case, present)
        test_cases.append(case)

    rows_read = sum(len(r.rows) for r in sheet_results)
    duplicates = [c for c in test_cases if len(c["sources"]) > 1]
    checks = [chk for r in sheet_results for chk in r.checks]
    checks.append(
        {
            "name": "Every test row is in exactly one test case",
            "ok": sum(len(c["sources"]) for c in test_cases) == rows_read,
            "detail": f"{rows_read} test rows read, {sum(len(c['sources']) for c in test_cases)} referenced by test cases",
        }
    )
    checks.append(
        {
            "name": "Test ids are unique",
            "ok": len({c["test_id"] for c in test_cases}) == len(test_cases),
            "detail": f"{len(test_cases)} test cases",
        }
    )
    skipped = [s for r in sheet_results for s in r.skipped]

    return {
        "sheets": [r.info for r in sheet_results],
        "report": {
            "sheets_total": len(sheet_results),
            "sheets_extracted": sum(1 for r in sheet_results if r.info.get("status") == "extracted"),
            "rows_read": rows_read,
            "unique_test_cases": len(test_cases),
            "duplicate_ids": len(duplicates),
            "identical_duplicates": sum(1 for c in duplicates if not c["conflicts"]),
            "conflicting_duplicates": sum(1 for c in duplicates if c["conflicts"]),
            "totals_rows": sum(len(r.info.get("totals_rows", [])) for r in sheet_results),
            "skipped_rows": skipped,
            "test_cases_with_warnings": sum(1 for c in test_cases if c["warnings"]),
            "checks": checks,
            "checks_passed": sum(1 for c in checks if c["ok"]),
            "checks_failed": sum(1 for c in checks if not c["ok"]),
        },
        "test_cases": test_cases,
    }
