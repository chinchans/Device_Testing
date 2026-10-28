"""Store an uploaded test suite workbook per operator stage and summarise its sheets."""

from __future__ import annotations

import hashlib
import json
import re
import shutil
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from core.config import get_settings
from observability.logging import logger

from .extract import extract_workbook

SUPPORTED_SUFFIXES = {".xlsx", ".xlsm"}
STAGE_ID_RE = re.compile(r"^[a-z0-9_-]{1,32}$")
META_FILE = "latest.json"
EXTRACTED_FILE = "extracted.json"


class SuiteUploadError(ValueError):
    """Bad request input (unknown stage, missing file, unsupported or unreadable workbook)."""


def _stage_dir(stage_id: str) -> Path:
    if not STAGE_ID_RE.match(stage_id):
        raise SuiteUploadError(f"Invalid stage id: {stage_id}")
    path = get_settings().resolve_path("data/test_suites") / stage_id
    path.mkdir(parents=True, exist_ok=True)
    return path


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for block in iter(lambda: fh.read(1 << 20), b""):
            digest.update(block)
    return digest.hexdigest()


def _cell_filled(value: Any) -> bool:
    return value is not None and not (isinstance(value, str) and not value.strip())


def _summarise_sheets(path: Path) -> list[dict[str, Any]]:
    """Count the rows and columns that actually hold data on every sheet."""
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError) as exc:
        raise SuiteUploadError(f"Could not open the workbook: {exc}") from exc
    try:
        sheets = []
        for ws in wb.worksheets:
            rows = 0
            columns = 0
            for values in ws.iter_rows(values_only=True):
                filled = [i for i, v in enumerate(values) if _cell_filled(v)]
                if filled:
                    rows += 1
                    columns = max(columns, filled[-1] + 1)
            sheets.append(
                {
                    "name": ws.title,
                    "visible": ws.sheet_state == "visible",
                    "rows": rows,
                    "columns": columns,
                }
            )
        return sheets
    finally:
        wb.close()


def upload(stage_id: str, file_path: str) -> dict[str, Any]:
    """Copy the workbook into the stage folder (replacing the previous one) and save its sheet summary."""
    source = Path(file_path).expanduser()
    if not source.is_file():
        raise SuiteUploadError(f"File not found: {file_path}")
    suffix = source.suffix.lower()
    if suffix == ".xls":
        raise SuiteUploadError("Legacy .xls workbooks are not supported. Save the file as .xlsx and upload it again.")
    if suffix not in SUPPORTED_SUFFIXES:
        raise SuiteUploadError(f"Unsupported file type {suffix or '(none)'}. Upload an Excel workbook (.xlsx).")

    sheets = _summarise_sheets(source)
    if not sheets:
        raise SuiteUploadError("The workbook has no sheets")

    folder = _stage_dir(stage_id)
    for old in folder.iterdir():
        if old.is_file():
            old.unlink()
    stored = folder / source.name
    shutil.copy2(source, stored)

    visible = [s for s in sheets if s["visible"]]
    record = {
        "stage_id": stage_id,
        "file_name": source.name,
        "source_path": str(source),
        "stored_path": str(stored),
        "size": stored.stat().st_size,
        "sha256": _sha256(stored),
        "uploaded_at": datetime.now().isoformat(timespec="seconds"),
        "sheets": sheets,
        "summary": {
            "sheets": len(sheets),
            "visible_sheets": len(visible),
            "hidden_sheets": len(sheets) - len(visible),
            "rows": sum(s["rows"] for s in visible),
        },
    }
    (folder / META_FILE).write_text(json.dumps(record, indent=2), encoding="utf-8")
    return record


def load_saved(stage_id: str) -> dict[str, Any] | None:
    path = _stage_dir(stage_id) / META_FILE
    if not path.is_file():
        return None
    try:
        record = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        logger.warning("Ignoring unreadable test suite record for %s: %s", stage_id, exc)
        return None
    if not Path(record.get("stored_path", "")).is_file():
        return None
    return record


def extraction_summary(extracted: dict[str, Any]) -> dict[str, Any]:
    """Everything except the test case blocks, for the UI."""
    return {k: v for k, v in extracted.items() if k != "test_cases"}


def extract(stage_id: str) -> dict[str, Any]:
    """Extract the stored workbook into extracted.json (one block per unique test id)."""
    record = load_saved(stage_id)
    if record is None:
        raise SuiteUploadError("Upload a test suite workbook first")
    stored = Path(record["stored_path"])
    try:
        result = extract_workbook(stored)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError) as exc:
        raise SuiteUploadError(f"Could not read the workbook: {exc}") from exc

    output = _stage_dir(stage_id) / EXTRACTED_FILE
    extracted = {
        "stage_id": stage_id,
        "source": {
            "file_name": record["file_name"],
            "sha256": record["sha256"],
            "uploaded_at": record["uploaded_at"],
        },
        "extracted_at": datetime.now().isoformat(timespec="seconds"),
        "output_path": str(output),
        **result,
    }
    output.write_text(json.dumps(extracted, indent=2, ensure_ascii=False), encoding="utf-8")
    return extracted


def load_extracted(stage_id: str) -> dict[str, Any] | None:
    path = _stage_dir(stage_id) / EXTRACTED_FILE
    if not path.is_file():
        return None
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        logger.warning("Ignoring unreadable extraction for %s: %s", stage_id, exc)
        return None


def clear_saved(stage_id: str) -> None:
    folder = _stage_dir(stage_id)
    for path in folder.iterdir():
        if path.is_file():
            path.unlink()
